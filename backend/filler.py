from docx import Document
from io import BytesIO
import os
import copy

def merge_documents(doc1, doc2):
    """
    Appends contents of doc2 to the end of doc1.
    Uses deepcopy so elements are copied, not moved out of doc2.
    """
    doc1.add_page_break()

    for element in doc2.element.body:
        # Avoid SectPr (section properties) which can break page layout if copied blindly
        if element.tag.endswith("sectPr"):
            continue
        doc1.element.body.append(copy.deepcopy(element))

    return doc1


def fill_template(template_path, data):
    doc = Document(template_path)

    # Generic placeholder replacement for paragraphs
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            placeholder = "{{" + key + "}}"
            if placeholder in paragraph.text:
                # Handle nested dicts from AI if they occur
                if isinstance(value, dict) and len(value) == 1:
                    value = list(value.values())[0]

                # If value is a list, join it with newlines
                if isinstance(value, list):
                    value = "\n".join([str(v) for v in value])

                # Handle multi-line value
                if "\n" in str(value):
                    lines = str(value).split("\n")
                    paragraph.text = paragraph.text.replace(placeholder, lines[0])
                    for line in lines[1:]:
                        paragraph.add_run("\n" + line)
                else:
                    paragraph.text = paragraph.text.replace(placeholder, str(value))

    # Support for the {{CONTENT}} placeholder (Full Content)
    if "FULL_CONTENT" in data:
        for paragraph in doc.paragraphs:
            if "{{CONTENT}}" in paragraph.text:
                paragraph.text = paragraph.text.replace("{{CONTENT}}", data["FULL_CONTENT"])
                break

    return doc


def clean_for_pdf(text):
    if not text:
        return ""
    # Mapping of common unicode characters to latin-1 equivalents or safe characters
    replacements = {
        "\u2022": "-",  # Bullet
        "\u2023": "-",  # Triangular bullet
        "\u2043": "-",  # Hyphen bullet
        "\u2013": "-",  # En dash
        "\u2014": "-",  # Em dash
        "\u2018": "'",  # Left single quote
        "\u2019": "'",  # Right single quote
        "\u201c": '"',  # Left double quote
        "\u201d": '"',  # Right double quote
        "\uf0b7": "-",  # Wingdings bullet
        "\uf02d": "-",  # Wingdings hyphen
        "\u25cf": "-",  # Large circle bullet
        "\u2026": "...",  # Ellipsis
        "\u00a0": " ",  # Non-breaking space
    }
    for k, v in replacements.items():
        text = text.replace(k, v)

    # Aggressive cleanup: replace anything non-latin-1 with a safe character
    cleaned = ""
    for char in text:
        try:
            char.encode("latin-1")
            cleaned += char
        except UnicodeEncodeError:
            cleaned += "-"

    return cleaned


def generate_recommendation_pdf(data):
    from fpdf import FPDF
    import tempfile

    pdf = FPDF()
    pdf.add_page()

    # Add Logo if exists
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    logo_path = os.path.join(base_dir, "templates", "media", "media", "image2.jpg")

    if os.path.exists(logo_path):
        pdf.image(logo_path, x=10, y=8, w=190)
        pdf.ln(35)
    else:
        print(f"PDF Logo not found at: {logo_path}")

    # Header: "Comments from Hexaview (Name)"
    pdf.set_font("Arial", "BU", 14)
    pdf.set_text_color(0, 51, 153)
    name = clean_for_pdf(data.get("EVALUATOR", "Candidate"))
    pdf.cell(0, 10, f"Comments from Hexaview ({name})", ln=True, align="L")
    pdf.ln(5)

    sections = [
        ("Summary", "REC_SUMMARY"),
        ("Education", "REC_EDUCATION"),
        ("Employer and Work", "REC_WORK"),
        ("Candidate Strengths", "REC_STRENGTHS"),
        ("Hexaview Hiring Recommendation", "REC_RECOMMENDATION"),
    ]

    for title, key in sections:
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Arial", "B", 11)
        pdf.cell(0, 8, title, ln=True)

        pdf.set_font("Arial", size=10)
        value = data.get(key, "Not specified")

        # Robust extraction: handle dicts, lists, and strings
        if isinstance(value, dict):
            if len(value) == 1:
                value = list(value.values())[0]
            else:
                value = str(value)

        if isinstance(value, list):
            content = "\n".join([str(v) for v in value])
        else:
            content = str(value)

        content = clean_for_pdf(content)

        pdf.multi_cell(0, 5, content)
        pdf.ln(4)

        if title == "Summary":
            pdf.set_draw_color(200, 200, 200)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(4)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        pdf.output(tmp.name)
        with open(tmp.name, "rb") as f:
            pdf_bytes = f.read()
    os.unlink(tmp.name)
    return pdf_bytes


def merge_pdfs(rec_pdf_bytes, original_pdf_bytes):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()

    # Add newly generated recommendation page
    rec_reader = PdfReader(BytesIO(rec_pdf_bytes))
    writer.add_page(rec_reader.pages[0])

    # Add original pages, but skip the first page if it looks like a recommendation page
    orig_reader = PdfReader(BytesIO(original_pdf_bytes))

    start_page = 0
    if len(orig_reader.pages) > 0:
        first_page_text = orig_reader.pages[0].extract_text()
        markers = ["Comments from Hexaview", "Hiring Recommendation", "Candidate Strengths"]
        if any(m in first_page_text for m in markers):
            print("Detected existing recommendation page in PDF, skipping it.")
            start_page = 1

    for i in range(start_page, len(orig_reader.pages)):
        writer.add_page(orig_reader.pages[i])

    out_io = BytesIO()
    writer.write(out_io)
    out_io.seek(0)
    return out_io


def copy_doc_elements(source_bytes, target_doc):
    """
    Copies paragraphs and tables from source doc bytes to target doc.
    Skips the first few elements if they look like a recommendation page.
    """
    source_doc = Document(BytesIO(source_bytes))

    skip_elements_count = 0
    found_rec_header = False
    for i, source_p in enumerate(source_doc.paragraphs[:10]):
        if "Comments from Hexaview" in source_p.text:
            found_rec_header = True
            break

    if found_rec_header:
        temp_skip_count = 0
        for element in source_doc.element.body:
            temp_skip_count += 1
            xml_text = "".join(element.xpath(".//w:t"))
            if "Hexaview Hiring Recommendation" in xml_text or element.tag.endswith("sectPr"):
                break
            if temp_skip_count > 50:
                break
        skip_elements_count = temp_skip_count

    for p in target_doc.paragraphs:
        if "{{CONTENT}}" in p.text:
            p.text = ""
            parent = p._element.getparent()

            for i, element in enumerate(source_doc.element.body):
                if i < skip_elements_count:
                    continue
                if element.tag.endswith("sectPr"):
                    continue
                # deepcopy is required: lxml append() MOVES nodes (removes from source)
                # Using deepcopy ensures the original source document is not drained
                parent.append(copy.deepcopy(element))

            parent.remove(p._element)
            break
    return target_doc


def docx_to_pdf_bytes(docx_bytes):
    """
    Convert a DOCX (as bytes) to PDF bytes using
    docx2pdf (Windows) or LibreOffice (Linux/Docker).
    Uses a unique per-request LibreOffice profile to avoid
    lock contention when multiple users hit the server simultaneously.
    """
    import tempfile
    import platform
    import subprocess

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(docx_bytes)
        docx_path = tmp.name

    pdf_path = docx_path.replace(".docx", ".pdf")

    try:
        if platform.system() == "Windows":
            from docx2pdf import convert
            convert(docx_path, pdf_path)
        else:
            # Unique profile dir prevents lock issues with concurrent users
            profile_dir = tempfile.mkdtemp(prefix="lo_profile_")
            try:
                result = subprocess.run(
                    [
                        "soffice",
                        "--headless",
                        "--invisible",
                        "--nodefault",
                        "--nofirststartwizard",
                        "--nolockcheck",
                        "--nologo",
                        "--norestore",
                        f"-env:UserInstallation=file://{profile_dir}",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        os.path.dirname(docx_path),
                        docx_path,
                    ],
                    capture_output=True,
                    text=True,
                    timeout=120,
                )
            finally:
                import shutil
                shutil.rmtree(profile_dir, ignore_errors=True)

            if result.returncode != 0:
                error_msg = (
                    f"LibreOffice conversion failed.\n"
                    f"STDOUT: {result.stdout}\n"
                    f"STDERR: {result.stderr}"
                )
                print(error_msg)
                raise RuntimeError(error_msg)

        with open(pdf_path, "rb") as f:
            return f.read()
    finally:
        if os.path.exists(docx_path):
            os.unlink(docx_path)
        if os.path.exists(pdf_path):
            os.unlink(pdf_path)


def generate_multi_page_resume(rec_template_path, brand_template_path, data, original_bytes, filename):
    """
    Produces a PDF with:
      - Page 1   : Hexaview recommendation / evaluation front page
      - Page 2+  : The original resume, completely unchanged
    """

    # ── PDF input ──────────────────────────────────────────────────────────────
    if filename.lower().endswith(".pdf"):
        rec_pdf = generate_recommendation_pdf(data)
        return merge_pdfs(rec_pdf, original_bytes)

    # ── DOCX input ─────────────────────────────────────────────────────────────
    # The key insight: never merge at the DOCX level, because LibreOffice
    # re-renders combined DOCX files and can change fonts/layout of the original.
    # Instead:
    #   1. Generate the recommendation front page as a clean PDF (via fpdf).
    #   2. Convert the ORIGINAL, untouched DOCX to PDF (LibreOffice only
    #      ever sees the file the candidate submitted - nothing is modified).
    #   3. Merge the two PDFs.
    if filename.lower().endswith(".docx"):
        # Step 1: recommendation front page -> PDF bytes
        rec_pdf_bytes = generate_recommendation_pdf(data)

        # Step 2: original DOCX -> PDF bytes (original file passed as-is)
        original_pdf_bytes = docx_to_pdf_bytes(original_bytes)

        # Step 3: merge rec PDF (front) + original resume PDF (rest)
        return merge_pdfs(rec_pdf_bytes, original_pdf_bytes)

    # ── Fallback (unsupported format) ──────────────────────────────────────────
    raise ValueError(f"Unsupported file format: {filename}")
