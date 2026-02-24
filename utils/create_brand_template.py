from docx import Document
import os

def identify_and_keep_branding(doc):
    # Based on scanning, Index 0 is the logo banner.
    indices_to_keep = {0}
    
    print(f"Keeping paragraphs: {indices_to_keep}")
    
    # Delete all other paragraphs
    for i in range(len(doc.paragraphs) - 1, -1, -1):
        if i not in indices_to_keep:
            p = doc.paragraphs[i]._element
            p.getparent().remove(p)
            
    # Delete all tables
    for table in doc.tables:
        t = table._element
        t.getparent().remove(t)
        
    # Add {{CONTENT}} for the resume part
    doc.add_paragraph("{{CONTENT}}")

if __name__ == "__main__":
    branding_path = os.path.join("templates", "Resume - Piyush Nirwan.docx")
    output_path = os.path.join("templates", "branded_template_final.docx")
    rec_template_path = os.path.join("templates", "recommendation_template.docx")
    
    if os.path.exists(branding_path):
        doc = Document(branding_path)
        identify_and_keep_branding(doc)
        doc.save(output_path)
        print(f"Branding Template saved to {output_path}")

        
        # Also trigger recommendation template creation to be sure
        from create_recommendation_template import create_recommendation_template
        create_recommendation_template(output_path, rec_template_path)
    else:
        print(f"Error: {branding_path} not found.")
